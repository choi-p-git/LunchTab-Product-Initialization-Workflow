from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

CSV_ENCODINGS = (
    "utf-8-sig",
    "utf-8",
    "utf-16",
    "utf-16-le",
    "utf-16-be",
    "cp1252",
    "iso-8859-1",
)


def _decode_csv_bytes(data: bytes, path: Path) -> str:
    errors = []
    for encoding in CSV_ENCODINGS:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError as error:
            errors.append(f"{encoding}: byte {error.start}")
            continue
        if _looks_like_csv_text(text):
            return text
        errors.append(f"{encoding}: decoded text did not look like CSV")
    raise UnicodeDecodeError(
        "csv",
        data,
        0,
        1,
        f"could not decode {path} using supported encodings ({'; '.join(errors)})",
    )


def _looks_like_csv_text(text: str) -> bool:
    sample = text[:4096]
    if not sample.strip():
        return True
    if "\x00" in sample:
        return False
    first_line = sample.splitlines()[0] if sample.splitlines() else sample
    return any(delimiter in first_line for delimiter in (",", "\t", ";"))


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    data = path.read_bytes()
    text = _decode_csv_bytes(data, path)
    reader = csv.DictReader(io.StringIO(text, newline=""), strict=True)
    if reader.fieldnames is None:
        raise ValueError(f"CSV has no header row: {path}")
    return list(reader.fieldnames), list(reader)


def write_csv(path: Path, fieldnames: Iterable[str], rows: Iterable[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(fieldnames), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def read_inventory_workbook(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            for raw_row in rows:
                headers = ["" if value is None else str(value).strip() for value in raw_row]
                if {"Stock", "Item", "Retail Price", "Department", "Barcode"}.issubset(headers):
                    records: list[dict[str, str]] = []
                    for data_row in rows:
                        values = ["" if value is None else str(value).strip() for value in data_row]
                        if not any(values):
                            continue
                        padded = values + [""] * (len(headers) - len(values))
                        records.append(dict(zip(headers, padded, strict=False)))
                    return headers, records
    finally:
        workbook.close()
    raise ValueError(f"Could not find Odin inventory headers in {path}")
